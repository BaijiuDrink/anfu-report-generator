import datetime
from types import SimpleNamespace

from docx import Document
from docx.oxml.ns import qn
from PIL import Image
import pytest

from services.report_service import ReportService, ReportValidationError


def test_cover_date_does_not_depend_on_locale_encoding(monkeypatch):
    import report_builder

    class AsciiOnlyDate:
        year = 2026
        month = 10
        day = 1

        def strftime(self, pattern):
            pattern.encode("ascii")
            return datetime.datetime(self.year, self.month, self.day).strftime(pattern)

    class FixedClock:
        @staticmethod
        def now():
            return AsciiOnlyDate()

    monkeypatch.setattr(
        report_builder, "datetime", SimpleNamespace(datetime=FixedClock)
    )
    document = report_builder.ReportBuilder("示例项目").document
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert "报告编号：PT-20261001-001" in text
    assert "测试日期：2026年10月01日" in text
    assert "报告日期：2026年10月01日" in text


def test_empty_findings_are_rejected():
    with pytest.raises(ReportValidationError, match="请先录入至少一个漏洞"):
        ReportService().check([])


def test_missing_images_are_reported_but_marker_is_preserved(tmp_path):
    finding = {"name": "x", "verify_steps": "[截图: missing.png]"}
    check = ReportService().check([finding], base_dir=tmp_path)
    assert check.missing_images == [tmp_path / "missing.png"]
    assert finding["verify_steps"] == "[截图: missing.png]"


def test_generate_calls_existing_builder_in_order(tmp_path, monkeypatch):
    calls = []

    class FakeBuilder:
        def __init__(self, project_name):
            calls.append(("init", project_name))

        def add_summary_section(self, findings):
            calls.append(("summary", findings))

        def add_findings_section(self, findings):
            calls.append(("details", findings))

        def save(self, path):
            calls.append(("save", path))

    monkeypatch.setattr("services.report_service.ReportBuilder", FakeBuilder)
    findings = [{"name": "x", "verify_steps": ""}]
    output = tmp_path / "report.docx"
    ReportService().generate("demo", findings, output)
    assert [item[0] for item in calls] == ["init", "summary", "details", "save"]


def test_real_docx_preserves_fields_and_embeds_evidence(tmp_path):
    image_path = tmp_path / "proof.png"
    Image.new("RGB", (800, 400), "red").save(image_path)
    finding = {
        "name": "OAuth凭证泄露",
        "url": "https://example.test/app",
        "network_zone": "互联网",
        "risk_level": "高危",
        "description": "生产环境客户端凭证硬编码",
        "verify_steps": f"检查前端资源\n[截图: {image_path}]",
        "verify_result": "已确认",
        "fix_suggestion": "吊销并迁移到后端",
        "fix_priority": "紧急",
        "fix_verify": "重新扫描",
    }
    output = tmp_path / "report.docx"

    ReportService().generate("示例项目", [finding], output)

    document = Document(output)
    text = "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [
            paragraph.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
            for paragraph in cell.paragraphs
        ]
    )
    assert "OAuth凭证泄露" in text
    assert "https://example.test/app" in text
    assert "生产环境客户端凭证硬编码" in text
    assert len(document.inline_shapes) == 1
    assert any("/image" in rel.reltype for rel in document.part.rels.values())


def test_poc_exp_is_exported_as_shaded_monospaced_code_block(tmp_path):
    code = "curl -X POST /test\nif authorized:\n    run_exp()"
    output = tmp_path / "report.docx"

    ReportService().generate(
        "示例项目",
        [{"name": "命令执行", "poc_exp": code}],
        output,
    )

    document = Document(output)
    code_cells = [
        table.cell(0, 0) for table in document.tables if table.cell(0, 0).text == code
    ]

    assert code_cells
    cell = code_cells[0]
    shading = cell._tc.tcPr.find(qn("w:shd"))
    assert shading is not None
    assert shading.get(qn("w:fill")) == "F3F4F6"
    assert all(
        run.font.name == "Consolas"
        for paragraph in cell.paragraphs
        for run in paragraph.runs
    )


@pytest.mark.parametrize("poc_exp", ["", None])
def test_empty_poc_exp_is_omitted_from_report(tmp_path, poc_exp):
    output = tmp_path / "report.docx"

    ReportService().generate(
        "示例项目",
        [{"name": "命令执行", "poc_exp": poc_exp}],
        output,
    )

    document = Document(output)

    assert all(paragraph.text != "POC / EXP" for paragraph in document.paragraphs)
    assert not document.tables
